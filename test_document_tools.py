"""Document-reading tests for the MarkItDown-backed read_document atom."""

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))

from winbrow.document_tools import (                                  # noqa: E402
    DOCUMENT_IMPL,
    get_document_tools,
    read_document,
)
from winbrow.registry import ToolRegistry                             # noqa: E402
from winbrow.utils.history import format_history                      # noqa: E402

try:
    import markitdown  # noqa: F401
    HAS_MARKITDOWN = True
except Exception:
    HAS_MARKITDOWN = False


def _write_pdf(path: Path) -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 18 Tf 72 700 Td (Hello PDF document) Tj ET"
    objects[3] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += b"%010d 00000 n \n" % offset
    out += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref)
    )
    path.write_bytes(bytes(out))


def _write_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "word/document.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:r><w:t>Dear reader, this is a DOCX test.</w:t></w:r></w:p></w:body>
</w:document>""",
        )


class DocumentToolsTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(prefix="winbrow_docs_")
        self.root = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_reads_text_formats(self):
        cases = {
            "note.md": ("# Title\n\nHello **world**.\n", "Hello **world**"),
            "note.txt": ("Plain text line one.\nLine two.\n", "Line two"),
            "data.csv": ("name,score\nAlice,10\nBob,20\n", "Alice"),
            "data.json": (json.dumps({"a": 1, "b": [2, 3]}), '"a": 1'),
            "page.html": ("<html><body><h1>Hi</h1><p>There</p></body></html>", "# Hi"),
        }
        for name, (content, needle) in cases.items():
            path = self.root / name
            path.write_text(content, encoding="utf-8")
            result = read_document(str(path))
            self.assertTrue(result["success"], result["output"])
            self.assertIn(needle, result["output"])
            self.assertEqual(result["path"], str(path.resolve()))

    @unittest.skipUnless(HAS_MARKITDOWN, "markitdown is not installed")
    def test_reads_office_and_pdf_formats(self):
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "score"])
        sheet.append(["Alice", 10])
        workbook.save(self.root / "data.xlsx")

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Deck Title"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
        box.text_frame.text = "Slide body text"
        deck.save(self.root / "deck.pptx")

        _write_docx(self.root / "letter.docx")
        _write_pdf(self.root / "doc.pdf")

        expectations = {
            "data.xlsx": "Alice",
            "deck.pptx": "Deck Title",
            "letter.docx": "DOCX test",
            "doc.pdf": "Hello PDF document",
        }
        for name, needle in expectations.items():
            result = read_document(str(self.root / name))
            self.assertTrue(result["success"], result["output"])
            self.assertIn(needle, result["output"])

    def test_pagination_returns_next_offset(self):
        path = self.root / "long.txt"
        body = "".join(f"line {i:04d}\n" for i in range(300))
        path.write_text(body, encoding="utf-8")

        first = read_document(str(path), max_chars=1000)
        self.assertTrue(first["success"], first["output"])
        self.assertEqual(first["offset"], 0)
        self.assertEqual(first["next_offset"], 1000)
        self.assertIn("TRUNCATED", first["output"])

        second = read_document(str(path), max_chars=1000, offset=first["next_offset"])
        self.assertTrue(second["success"], second["output"])
        self.assertEqual(second["offset"], 1000)
        self.assertIn("line 0100", second["output"])

    def test_document_output_survives_history_compression(self):
        marker = "FULL-DOCUMENT-TEXT"
        body = "x" * 12000 + marker + "y" * 12000
        history = [{
            "step": {"skill": "read_document", "args": {"path": "deck.pptx"}},
            "result": {"success": True, "output": body},
        }]
        block = format_history(history)
        self.assertIn(marker, block)

    def test_history_summarizes_large_write_args(self):
        history = [{
            "step": {
                "skill": "mcp_filesystem_write_file",
                "args": {"path": "summary.md", "content": "Q" * 25000},
            },
            "result": {"success": True, "output": "wrote 25000 chars"},
        }]
        block = format_history(history)
        self.assertIn("content=<25000 chars>", block)
        self.assertNotIn("Q" * 200, block)

    def test_max_chars_boundary_and_pagination(self):
        from winbrow import document_tools as dt
        self.assertEqual(dt.DEFAULT_MAX_CHARS, 60000)
        self.assertEqual(dt.MAX_OUTPUT_CHARS, 100000)
        path = self.root / "big.txt"
        path.write_text("x" * 120_000, encoding="utf-8")
        result = read_document(str(path), max_chars=100_001)
        self.assertTrue(result["success"], result["output"])
        self.assertEqual(result["next_offset"], 100_000)
        self.assertIn("offset=100000", result["output"])
        result = read_document(str(path), max_chars=500)
        self.assertEqual(result["next_offset"], 1000)

    def test_document_history_budget_is_newest_first(self):
        history = []
        for index in range(5):
            history.append({
                "step": {"skill": "read_document", "args": {"offset": index}},
                "result": {
                    "success": True,
                    "output": "z" * 20000 + f"marker-{index}",
                },
            })
        block = format_history(history)
        self.assertIn("marker-4", block)
        self.assertNotIn("marker-0", block)
        self.assertLess(len(block), 65000)

    def test_document_history_full_mode_keeps_all_steps(self):
        history = []
        for index in range(5):
            history.append({
                "step": {"skill": "read_document", "args": {"offset": index}},
                "result": {
                    "success": True,
                    "output": "z" * 20000 + f"marker-{index}",
                },
            })
        block = format_history(
            history, full=True, max_chars=200_000, run_id="br_doc")
        for index in range(5):
            self.assertIn(f"marker-{index}", block)
        self.assertGreater(len(block), 90_000)

    def test_missing_file_is_reported(self):
        result = read_document(str(self.root / "does-not-exist.pdf"))
        self.assertFalse(result["success"])
        self.assertIn("File not found", result["output"])

    def test_text_fallback_without_markitdown(self):
        path = self.root / "fallback.txt"
        path.write_text("fallback works", encoding="utf-8")
        with patch("winbrow.document_tools._get_markitdown", side_effect=ImportError):
            result = read_document(str(path))
        self.assertTrue(result["success"], result["output"])
        self.assertIn("fallback works", result["output"])

    def test_binary_fallback_without_markitdown_is_clear(self):
        path = self.root / "fallback.pdf"
        _write_pdf(path)
        with patch("winbrow.document_tools._get_markitdown", side_effect=ImportError):
            result = read_document(str(path))
        self.assertFalse(result["success"])
        self.assertIn("markitdown", result["output"])

    def test_tool_definition_and_registry(self):
        tools = get_document_tools()
        self.assertEqual(len(tools), 1)
        tool = tools[0]
        self.assertEqual(tool.name, "read_document")
        self.assertEqual(tool.impl, DOCUMENT_IMPL)
        self.assertEqual([arg.name for arg in tool.args], ["path", "max_chars", "offset"])

        registry = ToolRegistry()
        registry.register(tool)
        self.assertIs(registry.get("read_document"), tool)

    def test_skill_is_registered(self):
        from winbrow.brain.agent_loop import SKILL_REGISTRY
        skill = SKILL_REGISTRY.get("read_document")
        self.assertIsNotNone(skill)
        self.assertEqual(skill.skill, "read_document")
        self.assertIn("path", {arg.name for arg in skill.args})


if __name__ == "__main__":
    unittest.main()
