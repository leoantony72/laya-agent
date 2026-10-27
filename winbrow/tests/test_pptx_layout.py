"""Measured PPTX layout: aspect, auto-fit, tables, images, charts, validation."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from winbrow.skill_scripts import execute_skill_script_tool


def _run(spec: dict, timeout: int = 120) -> dict:
    result = execute_skill_script_tool(
        "run_skill_script",
        {"skill": "pptx", "script": "pptx_tool.py", "spec": spec, "timeout": timeout},
    )
    assert result.get("success"), result.get("output")
    return result


@pytest.fixture(scope="module")
def image(tmp_path_factory) -> Path:
    pytest.importorskip("PIL")
    from PIL import Image

    target = tmp_path_factory.mktemp("pptx-assets") / "photo.png"
    Image.new("RGB", (1400, 800), (217, 119, 87)).save(target)
    return target


def _deck_spec(output: Path, image: Path, aspect: str = "16:9") -> dict:
    return {
        "operation": "create",
        "output_path": str(output),
        "title": "Q3 Review",
        "subtitle": "Growth, risks, next steps",
        "author": "WinBrow",
        "theme": "midnight",
        "aspect": aspect,
        "footer": "Confidential",
        "validate": True,
        "slides": [
            {"type": "title", "title": "Q3 Review", "subtitle": "Growth, risks"},
            {"type": "bullets", "title": "What changed",
             "bullets": ["Revenue grew 18%", "Enterprise drove 70%", "Churn fell"],
             "notes": "Pause here."},
            {"type": "two_column", "title": "Wins and risks",
             "left": ["Enterprise expansion", "New partners"],
             "right": ["Supply chain", "Hiring lag"],
             "left_title": "Wins", "right_title": "Risks"},
            {"type": "table", "title": "Revenue by region",
             "headers": ["Region", "Revenue", "Growth"],
             "rows": [["North", "$4.2M", "18%"], ["South", "$3.1M", "12%"]]},
            {"type": "image", "title": "Product",
             "images": [{"path": str(image), "caption": "Product surface",
                         "credit": "WinBrow sample"}]},
            {"type": "chart", "title": "Growth", "chart_type": "bar",
             "categories": ["Q1", "Q2", "Q3"],
             "series": [{"name": "Revenue", "values": [2.1, 3.2, 4.2]}]},
            {"type": "quote", "quote": "Simplicity is the ultimate sophistication.",
             "attribution": "Leonardo da Vinci"},
            {"type": "diagram", "title": "Rollout", "steps": ["Pilot", "Harden", "Launch"]},
            {"type": "closing", "title": "Thank you", "subtitle": "Questions?"},
        ],
    }


def test_create_valid_deck_and_preview_every_slide(tmp_path, image):
    output = tmp_path / "deck.pptx"
    result = _run(_deck_spec(output, image))
    data = result["data"]
    assert data["slides"] == 9
    assert data["slide_size"] == [13.333, 7.5]
    validation = data["validation"]
    assert validation["ok"], validation["errors"]
    assert validation["slides"] == 9
    assert len(validation["rendered"]) == 9
    assert all(path.endswith(".png") for path in validation["rendered"])

    from pptx import Presentation

    deck = Presentation(str(output))
    assert len(deck.slides) == 9
    assert abs(deck.slide_width - 12192000) < 1000  # 13.333 in
    assert deck.slide_height == 6858000
    # Notes survive on the bullets slide.
    notes = [
        slide.notes_slide.notes_text_frame.text
        for slide in deck.slides if slide.has_notes_slide
    ]
    assert "Pause here." in notes


def test_four_by_three_dimensions(tmp_path, image):
    output = tmp_path / "deck-4x3.pptx"
    result = _run(_deck_spec(output, image, aspect="4:3"))
    assert result["data"]["slide_size"] == [10.0, 7.5]
    assert result["data"]["validation"]["ok"]


def test_long_bullets_split_into_continuation_slides(tmp_path):
    output = tmp_path / "long.pptx"
    bullets = [
        f"Point {index}: " + ("A long explanatory sentence that needs wrapping. " * 3)
        for index in range(1, 26)
    ]
    result = _run({
        "operation": "create",
        "output_path": str(output),
        "title": "Long",
        "theme": "corporate",
        "slides": [{"type": "bullets", "title": "Many bullets", "bullets": bullets}],
    })
    assert result["data"]["slides"] > 1
    assert result["data"]["validation"]["ok"]


def test_image_caption_and_credit_are_present(tmp_path, image):
    output = tmp_path / "image.pptx"
    _run({
        "operation": "create",
        "output_path": str(output),
        "title": "Image",
        "slides": [{"type": "image", "title": "Photo",
                    "images": [{"path": str(image), "caption": "The caption",
                                "credit": "The credit"}]}],
    })
    from pptx import Presentation

    texts = [
        shape.text
        for slide in Presentation(str(output)).slides
        for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    ]
    assert any("The caption" in text for text in texts)
    assert any("The credit" in text for text in texts)


def test_large_table_splits_and_stays_inside_slide(tmp_path):
    rows = [
        [f"Row {index}", "A long description " * 8, f"${index}.2M"]
        for index in range(1, 28)
    ]
    output = tmp_path / "big-table.pptx"
    result = _run({
        "operation": "create",
        "output_path": str(output),
        "title": "Big table",
        "theme": "corporate",
        "slides": [{"type": "table", "title": "Large table",
                    "headers": ["Name", "Description", "Value"],
                    "rows": rows}],
    })
    assert result["data"]["slides"] > 1
    assert result["data"]["validation"]["ok"]


def test_validate_operation(tmp_path, image):
    output = tmp_path / "check.pptx"
    _run(_deck_spec(output, image))
    result = _run({
        "operation": "validate",
        "path": str(output),
        "aspect": "16:9",
        "render_dir": str(tmp_path / "slides"),
    })
    assert result["data"]["ok"] is True
    assert result["data"]["slides"] == 9
    assert len(result["data"]["rendered"]) == 9


def test_missing_image_fails_without_placeholder(tmp_path):
    result = execute_skill_script_tool(
        "run_skill_script",
        {"skill": "pptx", "script": "pptx_tool.py", "spec": {
            "operation": "create",
            "output_path": str(tmp_path / "bad.pptx"),
            "slides": [{"type": "image", "title": "Missing",
                        "images": [{"path": str(tmp_path / "nope.png")}]}],
        }, "timeout": 90},
    )
    assert result["success"] is False
