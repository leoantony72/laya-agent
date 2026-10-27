"""
Measured PowerPoint layout engine (python-pptx).

Builds editable slides with a consistent grid, auto-fitted text, measured
tables, proportional images, native charts, and recurring chrome. Geometry is
validated after building; see ``winbrow.layout.validate.validate_pptx``.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

from .measure import fit_font_size, image_fit, line_height, text_width, wrap_text

log = logging.getLogger("winbrow.layout.pptx")

ASPECTS = {"16:9": (13.333, 7.5), "4:3": (10.0, 7.5), "16:10": (13.333, 8.333)}

THEMES: dict[str, dict[str, Any]] = {
    "midnight": {
        "font": "Segoe UI", "bg": "0B1220", "surface": "111C2E",
        "text": "F8FAFC", "muted": "9AA8BC", "accent": "38BDF8",
        "accent2": "A78BFA", "title": 34, "body": 17, "small": 11,
    },
    "corporate": {
        "font": "Segoe UI", "bg": "FFFFFF", "surface": "F3F6FB",
        "text": "111827", "muted": "5B6472", "accent": "1D4ED8",
        "accent2": "0EA5E9", "title": 31, "body": 17, "small": 11,
    },
    "editorial": {
        "font": "Georgia", "bg": "FAF9F5", "surface": "F1EEE7",
        "text": "141413", "muted": "6B6862", "accent": "D97757",
        "accent2": "6A9BCC", "title": 33, "body": 17, "small": 11,
    },
    "minimal": {
        "font": "Segoe UI", "bg": "FFFFFF", "surface": "F4F4F4",
        "text": "111111", "muted": "666666", "accent": "111111",
        "accent2": "888888", "title": 34, "body": 17, "small": 11,
    },
    "forest": {
        "font": "Segoe UI", "bg": "0F1F18", "surface": "16291F",
        "text": "F3F7F4", "muted": "A8BDB0", "accent": "4ADE80",
        "accent2": "FDE68A", "title": 33, "body": 17, "small": 11,
    },
}


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(str(hex_value).lstrip("#").upper())


@dataclass
class PptxTheme:
    name: str = "midnight"
    aspect: str = "16:9"
    footer: str = ""
    title: str = ""
    author: str = ""
    font: str | None = None
    title_size: int | None = None
    body_size: int | None = None
    colors: dict[str, str] = field(default_factory=dict)

    @property
    def palette(self) -> dict[str, Any]:
        base = dict(THEMES.get(self.name, THEMES["midnight"]))
        base.update(self.colors or {})
        if self.font:
            base["font"] = self.font
        if self.title_size:
            base["title"] = int(self.title_size)
        if self.body_size:
            base["body"] = int(self.body_size)
        return base

    def dimensions(self) -> tuple[float, float]:
        return ASPECTS.get(self.aspect, ASPECTS["16:9"])


class PptxBuilder:
    def __init__(self, theme: PptxTheme):
        self.theme = theme
        self.palette = theme.palette
        self.width, self.height = theme.dimensions()
        self.margin = 0.58
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.width)
        self.prs.slide_height = Inches(self.height)
        self.blank = self.prs.slide_layouts[6]
        self.content_left = self.margin
        self.content_right = self.width - self.margin
        self.content_top = 1.55
        self.content_bottom = self.height - 0.62
        self.content_width = self.content_right - self.content_left
        self.content_height = self.content_bottom - self.content_top
        if theme.title:
            self.prs.core_properties.title = theme.title
        if theme.author:
            self.prs.core_properties.author = theme.author

    # ------------------------------------------------------------------ atoms
    def _new_slide(self, *, footer: bool = True) -> Any:
        slide = self.prs.slides.add_slide(self.blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _rgb(self.palette["bg"])
        # Accent bar across the top.
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(self.width), Inches(0.055),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(self.palette["accent"])
        bar.line.fill.background()
        if footer:
            self._footer(slide)
        return slide

    def _footer(self, slide: Any) -> None:
        y = self.height - 0.46
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(self.margin), Inches(y),
            Inches(self.width - self.margin), Inches(y),
        )
        line.line.color.rgb = _rgb(self.palette["muted"])
        line.line.width = Pt(0.5)
        text = self.theme.footer or self.theme.title
        if text:
            self._text(
                slide, self.margin, y + 0.05, self.content_width * 0.75, 0.25,
                text, size=self.palette["small"], color=_rgb(self.palette["muted"]),
            )
        index = len(self.prs.slides)
        self._text(
            slide, self.width - self.margin - 0.9, y + 0.05, 0.9, 0.25,
            str(index), size=self.palette["small"],
            color=_rgb(self.palette["muted"]), align=PP_ALIGN.RIGHT,
        )

    def _text(
        self, slide: Any, x: float, y: float, w: float, h: float, text: str,
        *, size: float, color: RGBColor, bold: bool = False, italic: bool = False,
        align: Any = PP_ALIGN.LEFT, anchor: Any = MSO_ANCHOR.TOP,
        line_spacing: float = 1.12, font: str | None = None,
        margins: tuple[float, float, float, float] = (0, 0, 0, 0),
    ) -> Any:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = None
        frame.margin_left = Inches(margins[0])
        frame.margin_top = Inches(margins[1])
        frame.margin_right = Inches(margins[2])
        frame.margin_bottom = Inches(margins[3])
        frame.vertical_anchor = anchor
        lines = str(text or "").split("\n")
        first = True
        for line in lines:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.alignment = align
            paragraph.line_spacing = line_spacing
            run = paragraph.add_run()
            run.text = line
            run.font.name = font or self.palette["font"]
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
        return box

    @staticmethod
    def _bullet(paragraph: Any, level: int = 0) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        p_pr.set("marL", str(int(Inches(0.26 + max(0, level) * 0.24))))
        p_pr.set("indent", str(int(-Inches(0.2))))
        for child in list(p_pr):
            if child.tag in {qn("a:buChar"), qn("a:buNone"), qn("a:buAutoNum")}:
                p_pr.remove(child)
        bullet = parse_xml(f'<a:buChar {nsdecls("a")} char="•"/>')
        p_pr.append(bullet)

    def _bullets(
        self, slide: Any, x: float, y: float, w: float, h: float,
        bullets: Iterable[Any], *, size: float, color: RGBColor,
    ) -> Any:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = None
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        first = True
        for item in bullets:
            if isinstance(item, dict):
                text = str(item.get("text") or "")
                level = int(item.get("level") or 0)
            else:
                text, level = str(item), 0
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.level = min(4, max(0, level))
            paragraph.line_spacing = 1.08
            paragraph.space_after = Pt(7)
            self._bullet(paragraph, level)
            run = paragraph.add_run()
            run.text = text
            run.font.name = self.palette["font"]
            run.font.size = Pt(size)
            run.font.color.rgb = color
        return box

    def _frame(self, title: str, *, kicker: str = "") -> Any:
        slide = self._new_slide()
        y = self.margin
        if kicker:
            self._text(
                slide, self.content_left, y, self.content_width, 0.25,
                kicker.upper(), size=self.palette["small"], bold=True,
                color=_rgb(self.palette["accent"]),
            )
            y += 0.26
        size, _ = fit_font_size(
            title, self.content_width * 72, 1.0 * 72,
            max_size=self.palette["title"], min_size=22,
            spacing=1.05, style="sans",
        )
        self._text(
            slide, self.content_left, y, self.content_width, 0.95, title,
            size=size, bold=True, color=_rgb(self.palette["text"]),
            line_spacing=1.02,
        )
        return slide

    def _fit_bullets(self, bullets: list[Any], width: float, height: float,
                     *, max_size: int, min_size: int) -> int:
        size = max_size
        while size >= min_size:
            total = 0.0
            for item in bullets:
                text = str(item.get("text") if isinstance(item, dict) else item)
                lines = wrap_text(text, (width - 0.35) * 72, size)
                total += len(lines) * line_height(size, 1.08) + 7
            if total <= height * 72 * 0.92:
                return size
            size -= 1
        return min_size

    def _chunk_bullets(self, bullets: list[Any], width: float, height: float,
                       size: int) -> list[list[Any]]:
        chunks: list[list[Any]] = []
        current: list[Any] = []
        used = 0.0
        limit = height * 72 * 0.92
        for item in bullets:
            text = str(item.get("text") if isinstance(item, dict) else item)
            lines = wrap_text(text, (width - 0.35) * 72, size)
            needed = len(lines) * line_height(size, 1.08) + 7
            if current and used + needed > limit:
                chunks.append(current)
                current, used = [], 0.0
            current.append(item)
            used += needed
        if current:
            chunks.append(current)
        return chunks

    # ------------------------------------------------------------------ slides
    def add_title(self, title: str, subtitle: str = "", notes: str = "") -> Any:
        slide = self._new_slide(footer=False)
        # Bold side panel with the accent colour.
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(self.width * 0.73), 0,
            Inches(self.width * 0.27), Inches(self.height),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = _rgb(self.palette["accent"])
        panel.line.fill.background()
        self._text(
            slide, self.margin, 1.55, self.width * 0.62, 1.5, title,
            size=self.palette["title"] + 8, bold=True,
            color=_rgb(self.palette["text"]), line_spacing=0.98,
        )
        if subtitle:
            self._text(
                slide, self.margin, 3.25, self.width * 0.60, 1.3, subtitle,
                size=self.palette["body"] + 2, color=_rgb(self.palette["muted"]),
            )
        self._text(
            slide, self.margin, self.height - 0.95, self.width * 0.55, 0.35,
            self.theme.author or self.theme.title, size=self.palette["small"],
            color=_rgb(self.palette["muted"]),
        )
        self.add_notes(slide, notes)
        return slide

    def add_section(self, title: str, subtitle: str = "", notes: str = "") -> Any:
        slide = self._new_slide(footer=False)
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(self.height * 0.34),
            Inches(self.width), Inches(self.height * 0.32),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(self.palette["accent"])
        band.line.fill.background()
        self._text(
            slide, self.margin, self.height * 0.38, self.width - 2 * self.margin,
            self.height * 0.18, title, size=self.palette["title"] + 6,
            bold=True, color=_rgb(self.palette["bg"]), align=PP_ALIGN.LEFT,
        )
        if subtitle:
            self._text(
                slide, self.margin, self.height * 0.56,
                self.width - 2 * self.margin, 0.5, subtitle,
                size=self.palette["body"], color=_rgb(self.palette["bg"]),
            )
        self.add_notes(slide, notes)
        return slide

    def add_bullets(self, title: str, bullets: list[Any], notes: str = "",
                    kicker: str = "") -> list[Any]:
        if not bullets:
            bullets = [""]
        size = self._fit_bullets(
            bullets, self.content_width, self.content_height,
            max_size=self.palette["body"] + 6, min_size=13,
        )
        chunks = self._chunk_bullets(bullets, self.content_width, self.content_height, size)
        slides = []
        for index, chunk in enumerate(chunks):
            slide = self._frame(
                title + (" (continued)" if index else ""), kicker=kicker,
            )
            measured = 0.0
            for item in chunk:
                text = str(item.get("text") if isinstance(item, dict) else item)
                measured += len(wrap_text(text, (self.content_width - 0.35) * 72, size)) * line_height(size, 1.08) + 7
            offset = max(0.0, (self.content_height * 72 - measured) / 2 / 72)
            y = self.content_top + offset
            self._bullets(
                slide, self.content_left, y,
                self.content_width, self.content_bottom - y, chunk,
                size=size, color=_rgb(self.palette["text"]),
            )
            if index == 0:
                self.add_notes(slide, notes)
            slides.append(slide)
        return slides

    def add_two_column(self, title: str, left: list[Any], right: list[Any],
                       *, left_title: str = "", right_title: str = "",
                       notes: str = "", kicker: str = "") -> Any:
        slide = self._frame(title, kicker=kicker)
        gap = 0.35
        col_w = (self.content_width - gap) / 2
        sizes = [
            self._fit_bullets(items, col_w, self.content_height,
                              max_size=self.palette["body"], min_size=12)
            for items in (left, right)
        ]
        size = min(sizes)
        heights = []
        for items in (left, right):
            total = 0.0
            for item in items:
                text = str(item.get("text") if isinstance(item, dict) else item)
                total += len(wrap_text(text, (col_w - 0.35) * 72, size)) * line_height(size, 1.08) + 7
            heights.append(total / 72 + (0.38 if left_title or right_title else 0))
        offset = max(0.0, (self.content_height - max(heights)) / 2)
        for index, (items, label) in enumerate(((left, left_title), (right, right_title))):
            x = self.content_left + index * (col_w + gap)
            y = self.content_top + offset
            if label:
                self._text(
                    slide, x, y, col_w, 0.32, label, size=self.palette["body"],
                    bold=True, color=_rgb(self.palette["accent"]),
                )
                y += 0.38
            self._bullets(
                slide, x, y, col_w, self.content_bottom - y, items,
                size=size, color=_rgb(self.palette["text"]),
            )
        # A quiet vertical divider keeps the two columns aligned.
        divider = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(self.margin + col_w + gap / 2), Inches(self.content_top + offset),
            Inches(self.margin + col_w + gap / 2),
            Inches(self.content_top + offset + max(heights)),
        )
        divider.line.color.rgb = _rgb(self.palette["muted"])
        divider.line.width = Pt(0.5)
        self.add_notes(slide, notes)
        return slide

    def _table_widths(self, headers: list[str], rows: list[list[str]],
                      total_w: float, size: float) -> list[float]:
        columns = max([len(headers)] + [len(row) for row in rows]) if rows else len(headers)
        headers += [""] * (columns - len(headers))
        total_pt = total_w * 72
        min_pt = total_pt * 0.12
        max_pt = total_pt * 0.55
        natural = []
        for index in range(columns):
            values = [headers[index]] + [row[index] for row in rows if index < len(row)]
            widest = max((text_width(value, size) for value in values), default=40)
            natural.append(min(max(min_pt, widest + 0.22 * 72), max_pt))
        total = sum(natural)
        if total <= total_pt:
            extra = (total_pt - total) / max(1, columns)
            natural = [min(max_pt, value + extra) for value in natural]
        total = sum(natural) or 1.0
        return [value * total_pt / total for value in natural]

    def _row_height(self, row: list[str], widths: list[float], size: float) -> float:
        lines = 1
        for index, value in enumerate(row):
            width_pt = max(24.0, widths[index] - 0.16 * 72)
            lines = max(lines, len(wrap_text(value, width_pt, size)))
        return max(0.38, lines * line_height(size, 1.15) / 72 + 0.08)

    def add_table(self, title: str, headers: list[str], rows: list[list[str]],
                  *, notes: str = "", kicker: str = "") -> list[Any]:
        headers = [str(item) for item in headers]
        rows = [[str(cell) for cell in row] for row in rows]
        if not headers and rows:
            headers, rows = rows[0], rows[1:]
        size = 13
        widths = self._table_widths(headers, rows, self.content_width, size)
        header_h = self._row_height(headers, widths, size)
        limit = self.content_height - header_h - 0.1
        for candidate in (13, 12, 11, 10):
            candidate_widths = self._table_widths(headers, rows, self.content_width, candidate)
            candidate_header = self._row_height(headers, candidate_widths, candidate)
            candidate_limit = self.content_height - candidate_header - 0.1
            if all(
                self._row_height(row, candidate_widths, candidate) <= candidate_limit
                for row in rows
            ):
                size = candidate
                widths = candidate_widths
                header_h = candidate_header
                limit = candidate_limit
                break
        else:
            size = 10
            widths = self._table_widths(headers, rows, self.content_width, size)
            header_h = self._row_height(headers, widths, size)
            limit = self.content_height - header_h - 0.1

        # Split a row whose cell text is taller than a whole slide into several
        # continuation rows, so no cell is clipped or pushed off the slide.
        line_h_in = size * 1.15 / 72
        max_lines = max(1, int((limit - 0.08) / line_h_in))
        expanded: list[list[str]] = []
        for row in rows:
            if self._row_height(row, widths, size) <= limit:
                expanded.append(row)
                continue
            cell_lines = [
                wrap_text(row[index] if index < len(row) else "",
                          max(24.0, widths[index] - 0.16 * 72), size)
                for index in range(len(headers))
            ]
            while any(cell_lines):
                chunk = [lines[:max_lines] for lines in cell_lines]
                cell_lines = [lines[max_lines:] for lines in cell_lines]
                expanded.append([
                    "\n".join(part).strip() for part in chunk
                ])
        rows = expanded
        header_h = self._row_height(headers, widths, size)
        limit = self.content_height - header_h - 0.1

        chunks: list[list[list[str]]] = []
        current: list[list[str]] = []
        used = 0.0
        for row in rows:
            height = self._row_height(row, widths, size)
            if current and used + height > limit:
                chunks.append(current)
                current, used = [], 0.0
            current.append(row)
            used += height
        if current or not chunks:
            chunks.append(current)
        slides = []
        for index, chunk in enumerate(chunks):
            slide = self._frame(
                title + (" (continued)" if index else ""), kicker=kicker,
            )
            total_h = self._row_height(headers, widths, size)
            total_h += sum(self._row_height(row, widths, size) for row in chunk)
            y = self.content_top + max(0.0, (self.content_height - total_h) / 2)
            shape = slide.shapes.add_table(
                len(chunk) + 1, len(headers),
                Inches(self.content_left), Inches(y),
                Inches(self.content_width), Inches(total_h),
            )
            table = shape.table
            for col_index, width in enumerate(widths):
                table.columns[col_index].width = Inches(width / 72)
            for row_index, row in enumerate([headers] + chunk):
                table.rows[row_index].height = Inches(
                    self._row_height(row, widths, size) + 0.02
                )
                for col_index in range(len(headers)):
                    cell = table.cell(row_index, col_index)
                    value = row[col_index] if col_index < len(row) else ""
                    cell.text = value
                    cell.margin_left = Inches(0.08)
                    cell.margin_right = Inches(0.08)
                    cell.margin_top = Inches(0.03)
                    cell.margin_bottom = Inches(0.03)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    frame = cell.text_frame
                    frame.word_wrap = True
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = self.palette["font"]
                            run.font.size = Pt(size)
                            run.font.color.rgb = _rgb(
                                self.palette["bg"] if row_index == 0
                                else self.palette["text"]
                            )
                            if row_index == 0:
                                run.font.bold = True
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(
                        self.palette["accent"] if row_index == 0
                        else (self.palette["surface"] if row_index % 2 else self.palette["bg"])
                    )
            if index == 0:
                self.add_notes(slide, notes)
            slides.append(slide)
        return slides

    def add_images(self, title: str, images: list[dict], *, caption: str = "",
                   credit: str = "", notes: str = "", kicker: str = "") -> Any:
        slide = self._frame(title, kicker=kicker)
        images = [item for item in images if item.get("path")]
        if not images:
            raise ValueError("No valid image paths.")
        top = self.content_top
        caption_h = 0.62 if (caption or credit) else 0.0
        bottom = self.content_bottom - caption_h
        if len(images) == 1:
            item = images[0]
            with Image.open(item["path"]) as image:
                width, height = image.size
            box_w = self.content_width
            box_h = max(1.0, bottom - top)
            fit_w, fit_h, _, _ = image_fit(width, height, box_w * 72, box_h * 72)
            x = self.content_left + (box_w - fit_w / 72) / 2
            y = top + (box_h - fit_h / 72) / 2
            pic = slide.shapes.add_picture(
                item["path"], Inches(x), Inches(y),
                width=Inches(fit_w / 72), height=Inches(fit_h / 72),
            )
            pic.line.color.rgb = _rgb(self.palette["muted"])
            pic.line.width = Pt(0.5)
        else:
            columns = 2
            rows = (len(images) + columns - 1) // columns
            gap = 0.18
            cell_w = (self.content_width - gap) / columns
            cell_h = (bottom - top - gap * (rows - 1)) / max(1, rows)
            for index, item in enumerate(images[:columns * rows]):
                row, col = divmod(index, columns)
                with Image.open(item["path"]) as image:
                    width, height = image.size
                fit_w, fit_h, _, _ = image_fit(width, height, cell_w * 72, cell_h * 72)
                x = self.content_left + col * (cell_w + gap) + (cell_w - fit_w / 72) / 2
                y = top + row * (cell_h + gap) + (cell_h - fit_h / 72) / 2
                pic = slide.shapes.add_picture(
                    item["path"], Inches(x), Inches(y),
                    width=Inches(fit_w / 72), height=Inches(fit_h / 72),
                )
                pic.line.color.rgb = _rgb(self.palette["muted"])
                pic.line.width = Pt(0.5)
        if caption or credit:
            text = caption
            if credit:
                text = (text + "\n" if text else "") + credit
            self._text(
                slide, self.content_left, self.content_bottom - caption_h + 0.04,
                self.content_width, caption_h - 0.04, text, size=self.palette["small"],
                color=_rgb(self.palette["muted"]), align=PP_ALIGN.CENTER,
                line_spacing=1.12,
            )
        self.add_notes(slide, notes)
        return slide

    def add_chart(self, title: str, chart_type: str, categories: list[str],
                  series: list[dict], *, notes: str = "",
                  kicker: str = "") -> Any:
        slide = self._frame(title, kicker=kicker)
        data = CategoryChartData()
        data.categories = categories
        for item in series:
            data.add_series(str(item.get("name") or "Series"), item.get("values") or [])
        kind = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }.get(str(chart_type or "bar").lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)
        chart = slide.shapes.add_chart(
            kind, Inches(self.content_left), Inches(self.content_top),
            Inches(self.content_width), Inches(self.content_height), data,
        ).chart
        chart.has_title = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.font.name = self.palette["font"]
        chart.font.size = Pt(10)
        try:
            chart.font.color.rgb = _rgb(self.palette["text"])
        except Exception:
            pass
        try:
            chart.chart_area.format.fill.solid()
            chart.chart_area.format.fill.fore_color.rgb = _rgb(self.palette["surface"])
        except Exception:
            pass
        # python-pptx does not expose a reliable chart-area fill on every
        # version; write the chartSpace spPr explicitly.
        try:
            from pptx.oxml import parse_xml
            from pptx.oxml.ns import nsdecls, qn
            space = chart._chartSpace
            existing = space.find(qn("c:spPr"))
            if existing is not None:
                space.remove(existing)
            sp_pr = parse_xml(
                f'<c:spPr {nsdecls("c", "a")}>'
                f'<a:solidFill><a:srgbClr val="{self.palette["surface"]}"/></a:solidFill>'
                f'<a:ln><a:noFill/></a:ln>'
                f'</c:spPr>'
            )
            tx_pr = space.find(qn("c:txPr"))
            if tx_pr is not None:
                tx_pr.addprevious(sp_pr)
            else:
                space.append(sp_pr)
        except Exception:
            pass
        try:
            chart.category_axis.tick_labels.font.color.rgb = _rgb(self.palette["muted"])
            chart.value_axis.tick_labels.font.color.rgb = _rgb(self.palette["muted"])
        except Exception:
            pass
        palette = [
            self.palette["accent"], self.palette["accent2"],
            "#F59E0B", "#10B981", "#EF4444", "#8B5CF6",
        ]
        for index, item in enumerate(chart.series):
            color = _rgb(palette[index % len(palette)])
            if kind in {XL_CHART_TYPE.LINE_MARKERS}:
                item.format.line.color.rgb = color
                item.format.line.width = Pt(2.5)
            else:
                item.format.fill.solid()
                item.format.fill.fore_color.rgb = color
        plot = chart.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(9)
        labels.font.name = self.palette["font"]
        if kind not in {XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT}:
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
        else:
            labels.position = XL_LABEL_POSITION.BEST_FIT
        self.add_notes(slide, notes)
        return slide

    def add_quote(self, quote: str, attribution: str = "", notes: str = "") -> Any:
        slide = self._new_slide(footer=False)
        size = 30
        while size > 20:
            lines = wrap_text(f"“{quote}”", (self.width - 2 * self.margin - 0.6) * 72, size, style="sans")
            if len(lines) * line_height(size, 1.12) <= 4.0 * 72:
                break
            size -= 2
        lines = wrap_text(f"“{quote}”", (self.width - 2 * self.margin - 0.6) * 72, size, style="sans")
        quote_h = len(lines) * line_height(size, 1.12) / 72
        block_h = quote_h + (0.8 if attribution else 0)
        top = max(1.1, (self.height - block_h) / 2)
        mark = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches((self.width - 1.0) / 2), Inches(top - 0.35),
            Inches(1.0), Inches(0.055),
        )
        mark.fill.solid()
        mark.fill.fore_color.rgb = _rgb(self.palette["accent"])
        mark.line.fill.background()
        self._text(
            slide, self.margin, top,
            self.width - 2 * self.margin, quote_h + 0.1, f"“{quote}”",
            size=size, italic=True, color=_rgb(self.palette["text"]),
            line_spacing=1.12, align=PP_ALIGN.CENTER,
        )
        if attribution:
            self._text(
                slide, self.margin, top + quote_h + 0.25,
                self.width - 2 * self.margin, 0.4, f"— {attribution}",
                size=self.palette["body"], color=_rgb(self.palette["muted"]),
                align=PP_ALIGN.CENTER,
            )
        self.add_notes(slide, notes)
        return slide

    def add_closing(self, title: str = "Thank you", subtitle: str = "",
                    notes: str = "") -> Any:
        slide = self._new_slide(footer=False)
        self._text(
            slide, self.margin, 2.0, self.content_width, 1.2, title,
            size=self.palette["title"] + 8, bold=True,
            color=_rgb(self.palette["text"]), align=PP_ALIGN.CENTER,
        )
        if subtitle:
            self._text(
                slide, self.margin, 3.25, self.content_width, 0.8, subtitle,
                size=self.palette["body"], color=_rgb(self.palette["muted"]),
                align=PP_ALIGN.CENTER,
            )
        self._text(
            slide, self.margin, self.height - 1.0, self.content_width, 0.35,
            self.theme.author or self.theme.title, size=self.palette["small"],
            color=_rgb(self.palette["muted"]), align=PP_ALIGN.CENTER,
        )
        self.add_notes(slide, notes)
        return slide

    def add_diagram(self, title: str, steps: list[str], *, notes: str = "",
                    kicker: str = "") -> Any:
        slide = self._frame(title, kicker=kicker)
        steps = [str(step) for step in steps if str(step).strip()][:5]
        if not steps:
            raise ValueError("A diagram needs at least one step.")
        gap = 0.2
        box_w = (self.content_width - gap * (len(steps) - 1)) / len(steps)
        h = 1.2
        y = self.content_top + (self.content_height - h) / 2
        for index, step in enumerate(steps):
            x = self.content_left + index * (box_w + gap)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                Inches(box_w), Inches(h),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(self.palette["surface"])
            shape.line.color.rgb = _rgb(self.palette["accent"])
            shape.line.width = Pt(1.2)
            self._text(
                slide, x + 0.1, y + 0.1, box_w - 0.2, h - 0.2, step,
                size=13, color=_rgb(self.palette["text"]),
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            )
            if index < len(steps) - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(x + box_w + gap / 4), Inches(y + h / 2 - 0.12),
                    Inches(gap / 2), Inches(0.24),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = _rgb(self.palette["accent"])
                arrow.line.fill.background()
        self.add_notes(slide, notes)
        return slide

    def add_notes(self, slide: Any, notes: str) -> None:
        if not notes:
            return
        frame = slide.notes_slide.notes_text_frame
        frame.text = str(notes)

    def add_shape(self, kind: str, x: float, y: float, w: float, h: float,
                  text: str = "", *, fill: str | None = None,
                  line: str | None = None) -> Any:
        slide = self.prs.slides[-1]
        shape_type = {
            "rect": MSO_SHAPE.RECTANGLE,
            "rectangle": MSO_SHAPE.RECTANGLE,
            "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
            "oval": MSO_SHAPE.OVAL,
            "circle": MSO_SHAPE.OVAL,
            "arrow": MSO_SHAPE.RIGHT_ARROW,
        }.get(kind.lower(), MSO_SHAPE.RECTANGLE)
        shape = slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(w), Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill or self.palette["surface"])
        if line:
            shape.line.color.rgb = _rgb(line)
        else:
            shape.line.color.rgb = _rgb(self.palette["accent"])
        if text:
            shape.text_frame.text = text
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = self.palette["font"]
                    run.font.size = Pt(14)
                    run.font.color.rgb = _rgb(self.palette["text"])
        return shape

    def save(self, output: Path) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(output))


def build_deck(spec: dict) -> tuple[PptxBuilder, list[dict[str, Any]]]:
    """Build a deck from a slide spec; returns (builder, unresolved images)."""
    theme = PptxTheme(
        name=str(spec.get("theme") or "midnight").strip().lower(),
        aspect=str(spec.get("aspect") or "16:9").strip(),
        footer=str(spec.get("footer") or "").strip(),
        title=str(spec.get("title") or "").strip(),
        author=str(spec.get("author") or "").strip(),
        font=spec.get("font"),
        title_size=spec.get("title_size"),
        body_size=spec.get("body_size"),
        colors=dict(spec.get("colors") or {}),
    )
    builder = PptxBuilder(theme)
    slides = spec.get("slides") or []
    if not isinstance(slides, list) or not slides:
        slides = [{"type": "title", "title": theme.title or "Presentation"}]
    pending: list[dict[str, Any]] = []
    for index, raw in enumerate(slides):
        if isinstance(raw, str):
            raw = {"type": "bullets", "title": raw, "bullets": []}
        if not isinstance(raw, dict):
            continue
        kind = str(raw.get("type") or "").strip().lower()
        title = str(raw.get("title") or "").strip()
        notes = str(raw.get("notes") or "").strip()
        kicker = str(raw.get("kicker") or "").strip()
        if not kind:
            kind = "title" if index == 0 and not title else "bullets"
        if kind in {"title", "cover"}:
            builder.add_title(title or theme.title, str(raw.get("subtitle") or ""), notes)
        elif kind in {"section", "divider"}:
            builder.add_section(title, str(raw.get("subtitle") or ""), notes)
        elif kind in {"bullets", "content", "text"}:
            bullets = raw.get("bullets") or raw.get("content") or []
            if isinstance(bullets, str):
                bullets = [line.strip(" -") for line in bullets.splitlines() if line.strip()]
            builder.add_bullets(title, list(bullets), notes, kicker=kicker)
        elif kind in {"two_column", "columns"}:
            builder.add_two_column(
                title, list(raw.get("left") or []), list(raw.get("right") or []),
                left_title=str(raw.get("left_title") or ""),
                right_title=str(raw.get("right_title") or ""),
                notes=notes, kicker=kicker,
            )
        elif kind == "table":
            builder.add_table(
                title, list(raw.get("headers") or []),
                [list(row) for row in (raw.get("rows") or [])],
                notes=notes, kicker=kicker,
            )
        elif kind in {"image", "images"}:
            images = raw.get("images")
            if not images:
                images = [{"path": raw.get("image_path") or raw.get("path"),
                           "query": raw.get("query"),
                           "caption": raw.get("caption") or "",
                           "credit": raw.get("credit") or ""}]
            builder.add_images(
                title, [item for item in images if item.get("path")],
                caption=str(raw.get("caption") or (
                    images[0].get("caption") if len(images) == 1 else ""
                ) or ""),
                credit=str(raw.get("credit") or (
                    images[0].get("credit") if len(images) == 1 else ""
                ) or ""),
                notes=notes, kicker=kicker,
            )
            for item in images:
                if not item.get("path") and item.get("query"):
                    pending.append({"slide_index": len(builder.prs.slides),
                                    "item": item, "title": title, "raw": raw})
        elif kind == "chart":
            builder.add_chart(
                title, str(raw.get("chart_type") or "bar"),
                [str(item) for item in (raw.get("categories") or [])],
                list(raw.get("series") or []), notes=notes, kicker=kicker,
            )
        elif kind == "quote":
            builder.add_quote(
                str(raw.get("quote") or raw.get("text") or ""),
                str(raw.get("attribution") or ""), notes,
            )
        elif kind in {"closing", "end"}:
            builder.add_closing(title or "Thank you", str(raw.get("subtitle") or ""), notes)
        elif kind in {"diagram", "flow"}:
            builder.add_diagram(title, list(raw.get("steps") or []),
                                notes=notes, kicker=kicker)
        else:
            builder.add_bullets(title, [str(raw.get("text") or "")], notes, kicker=kicker)
    return builder, pending
