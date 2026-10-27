"""Automated and rendered validation for generated PDFs and PPTX decks."""

from __future__ import annotations

import logging
import shutil
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image

from .measure import wrap_text

log = logging.getLogger("winbrow.layout.validate")


def render_pdf_pages(pdf_path: str | Path, output_dir: str | Path,
                     scale: float = 1.5) -> list[str]:
    pdf_path = Path(pdf_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        import pypdfium2 as pdfium
    except Exception as exc:  # pragma: no cover
        log.warning("pypdfium2 unavailable: %s", exc)
        return []
    document = pdfium.PdfDocument(str(pdf_path))
    written: list[str] = []
    try:
        for index in range(len(document)):
            image = document[index].render(scale=scale).to_pil()
            target = output_dir / f"{pdf_path.stem}_{index + 1:02d}.png"
            image.save(target)
            written.append(str(target))
    finally:
        try:
            document.close()
        except Exception:
            pass
    return written


def _find_soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    for base in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(base).is_file():
            return base
    return None


def _render_pptx_with_soffice(pptx_path: Path, render_dir: Path) -> list[str]:
    soffice = _find_soffice()
    if not soffice:
        return []
    import subprocess

    render_dir.mkdir(parents=True, exist_ok=True)
    command = [
        soffice, "--headless", "--convert-to", "pdf",
        "--outdir", str(render_dir), str(pptx_path),
    ]
    try:
        subprocess.run(
            command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            timeout=120, check=False,
        )
    except Exception:
        return []
    pdf_path = render_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.is_file():
        return []
    return render_pdf_pages(pdf_path, render_dir)


def validate_pptx(
    pptx_path: str | Path,
    *,
    aspect: str | None = None,
    render_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Geometry, text-fit, image-ratio, and preview checks for a PPTX deck."""
    from pptx import Presentation

    from .preview import render_pptx_preview

    pptx_path = Path(pptx_path)
    errors: list[str] = []
    warnings: list[str] = []
    rendered: list[str] = []
    slide_count = 0
    size_inches = (0.0, 0.0)
    try:
        presentation = Presentation(str(pptx_path))
        slide_count = len(presentation.slides)
        slide_w = int(presentation.slide_width)
        slide_h = int(presentation.slide_height)
        size_inches = (round(slide_w / 914400, 3), round(slide_h / 914400, 3))
        if slide_count == 0:
            errors.append("Presentation has no slides.")
        if aspect:
            expected = {"16:9": (13.333, 7.5), "4:3": (10.0, 7.5),
                        "16:10": (13.333, 8.333)}.get(aspect)
            if expected and (
                abs(size_inches[0] - expected[0]) > 0.02
                or abs(size_inches[1] - expected[1]) > 0.02
            ):
                errors.append(
                    f"Slide size is {size_inches[0]}x{size_inches[1]}in; "
                    f"expected {expected[0]}x{expected[1]}in."
                )

        def bbox(shape):
            return (
                int(shape.left), int(shape.top),
                int(shape.left) + int(shape.width),
                int(shape.top) + int(shape.height),
            )

        def overlap(a, b):
            x0 = max(a[0], b[0])
            y0 = max(a[1], b[1])
            x1 = min(a[2], b[2])
            y1 = min(a[3], b[3])
            if x1 <= x0 or y1 <= y0:
                return 0.0
            return float((x1 - x0) * (y1 - y0))

        for index, slide in enumerate(presentation.slides, start=1):
            if not slide.shapes:
                errors.append(f"Slide {index}: no shapes.")
                continue
            content_shapes = []
            for shape in slide.shapes:
                try:
                    left, top, right, bottom = bbox(shape)
                except Exception:
                    continue
                if (
                    left < -1000 or top < -1000
                    or right > slide_w + 1000 or bottom > slide_h + 1000
                ):
                    errors.append(
                        f"Slide {index}: shape '{shape.name}' is outside the slide "
                        f"({left / 914400:.2f},{top / 914400:.2f}in)."
                    )
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    frame = shape.text_frame
                    inner_w = max(
                        1.0,
                        (shape.width - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700,
                    )
                    inner_h = max(
                        1.0,
                        (shape.height - (frame.margin_top or 0) - (frame.margin_bottom or 0)) / 12700,
                    )
                    needed = 0.0
                    for paragraph in frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs) or paragraph.text
                        size = 0.0
                        for run in paragraph.runs:
                            if run.font.size:
                                size = max(size, run.font.size.pt)
                        if size <= 0:
                            size = 18.0
                        lines = wrap_text(text, inner_w * 0.98, size)
                        spacing = float(paragraph.line_spacing or 1.12)
                        needed += len(lines) * size * spacing
                    if needed > inner_h + 3:
                        errors.append(
                            f"Slide {index}: text in '{shape.name}' needs "
                            f"{needed:.0f}pt but the box is {inner_h:.0f}pt high."
                        )
                    if top < slide_h - int(0.65 * 914400):
                        content_shapes.append(shape)
                elif getattr(shape, "has_chart", False):
                    content_shapes.append(shape)
                elif getattr(shape, "has_table", False):
                    table = shape.table
                    col_widths = [
                        max(10.0, column.width / 12700) for column in table.columns
                    ]
                    for row_index, row in enumerate(table.rows):
                        row_h = row.height / 12700
                        for col_index, cell in enumerate(row.cells):
                            text = cell.text.strip()
                            if not text:
                                continue
                            size = 0.0
                            for paragraph in cell.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    if run.font.size:
                                        size = max(size, run.font.size.pt)
                            if size <= 0:
                                size = 12.0
                            lines = wrap_text(
                                text,
                                max(10.0, col_widths[col_index] - 8),
                                size,
                            )
                            needed = len(lines) * size * 1.15
                            if needed > row_h + 5:
                                errors.append(
                                    f"Slide {index}: table cell '{text[:30]}' "
                                    f"needs {needed:.0f}pt but the row is "
                                    f"{row_h:.0f}pt high."
                                )
                    content_shapes.append(shape)
                elif "PICTURE" in str(shape.shape_type):
                    try:
                        with Image.open(BytesIO(shape.image.blob)) as image:
                            source_ratio = image.width / max(1, image.height)
                        box_ratio = shape.width / max(1, shape.height)
                        if abs(source_ratio - box_ratio) > 0.03 * max(1.0, source_ratio):
                            errors.append(
                                f"Slide {index}: image '{shape.name}' is distorted "
                                f"(source {source_ratio:.2f}, box {box_ratio:.2f})."
                            )
                    except Exception:
                        pass
                    content_shapes.append(shape)
            for first in range(len(content_shapes)):
                for second in range(first + 1, len(content_shapes)):
                    a = bbox(content_shapes[first])
                    b = bbox(content_shapes[second])
                    area_a = max(1, (a[2] - a[0]) * (a[3] - a[1]))
                    area_b = max(1, (b[2] - b[0]) * (b[3] - b[1]))
                    shared = overlap(a, b)
                    if shared > 0.15 * min(area_a, area_b):
                        warnings.append(
                            f"Slide {index}: '{content_shapes[first].name}' overlaps "
                            f"'{content_shapes[second].name}'."
                        )
        if render_dir:
            render_dir = Path(render_dir)
            rendered = _render_pptx_with_soffice(pptx_path, render_dir)
            if rendered:
                renderer = "libreoffice"
            else:
                rendered = render_pptx_preview(pptx_path, render_dir)
                renderer = "preview"
                warnings.append(
                    "No PowerPoint/LibreOffice renderer found; used the built-in "
                    "PIL preview. Chart drawing is simplified. Open the file in "
                    "PowerPoint for final confirmation."
                )
        else:
            renderer = ""
    except Exception as exc:
        errors.append(f"Could not inspect PPTX: {type(exc).__name__}: {exc}")
        renderer = ""

    return {
        "ok": not errors,
        "errors": errors[:40],
        "warnings": warnings[:40],
        "slides": slide_count,
        "slide_size": list(size_inches),
        "rendered": rendered,
        "renderer": renderer,
    }


def validate_pdf(
    pdf_path: str | Path,
    *,
    margins: tuple[float, float, float, float] = (56, 56, 56, 56),
    expected_size: tuple[float, float] | None = None,
    render_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Check text bounds, page geometry, blank pages, and render every page."""
    pdf_path = Path(pdf_path)
    errors: list[str] = []
    warnings: list[str] = []
    pages = 0
    page_size = (0.0, 0.0)
    try:
        import pdfplumber
    except Exception as exc:
        return {
            "ok": False,
            "errors": [f"pdfplumber unavailable: {exc}"],
            "warnings": [],
            "pages": 0,
            "rendered": [],
        }
    left_m, top_m, right_m, bottom_m = margins
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            pages = len(pdf.pages)
            if not pages:
                errors.append("PDF has no pages.")
            for index, page in enumerate(pdf.pages, start=1):
                width, height = float(page.width), float(page.height)
                if index == 1:
                    page_size = (width, height)
                    if expected_size and (
                        abs(width - expected_size[0]) > 1.5
                        or abs(height - expected_size[1]) > 1.5
                    ):
                        errors.append(
                            f"Page size is {width:.1f}x{height:.1f}pt; expected "
                            f"{expected_size[0]:.1f}x{expected_size[1]:.1f}pt."
                        )
                words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
                body_words = []
                for word in words:
                    x0, x1 = float(word["x0"]), float(word["x1"])
                    top, bottom = float(word["top"]), float(word["bottom"])
                    if x0 < -1 or x1 > width + 1 or top < -1 or bottom > height + 1:
                        errors.append(
                            f"Page {index}: text outside the page at "
                            f"({x0:.0f},{top:.0f})-({x1:.0f},{bottom:.0f})."
                        )
                        continue
                    # Header/footer live deliberately inside the margins.
                    if top >= top_m - 1 and bottom <= height - bottom_m + 1:
                        body_words.append(word)
                        if x0 < left_m - 2 or x1 > width - right_m + 2:
                            errors.append(
                                f"Page {index}: text crosses the content margin: "
                                f"'{word['text'][:40]}' at x={x0:.0f}-{x1:.0f}."
                            )
                if body_words:
                    body_left = min(float(w["x0"]) for w in body_words)
                    body_right = max(float(w["x1"]) for w in body_words)
                    usable = width - left_m - right_m
                    if usable > 0 and not page.images and (body_right - body_left) < usable * 0.55:
                        warnings.append(
                            f"Page {index}: body text uses only "
                            f"{(body_right - body_left) / usable:.0%} of the content width."
                        )
                else:
                    has_image = bool(page.images)
                    if not has_image:
                        warnings.append(f"Page {index}: no text or images (blank page).")
    except Exception as exc:
        errors.append(f"Could not inspect PDF: {type(exc).__name__}: {exc}")

    rendered: list[str] = []
    if render_dir:
        try:
            rendered = render_pdf_pages(pdf_path, render_dir)
            if len(rendered) != pages:
                warnings.append(
                    f"Rendered {len(rendered)} of {pages} pages for visual inspection."
                )
        except Exception as exc:
            warnings.append(f"Could not render pages: {type(exc).__name__}: {exc}")

    return {
        "ok": not errors,
        "errors": errors[:40],
        "warnings": warnings[:40],
        "pages": pages,
        "page_size": [round(page_size[0], 2), round(page_size[1], 2)],
        "rendered": rendered,
    }
